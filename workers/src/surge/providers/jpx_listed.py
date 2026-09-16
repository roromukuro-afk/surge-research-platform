"""JP security master from the JPX listed issue workbook (data_j.xlsx).

The workbook carries JPX's own market / product classification, which is what
universe-1.0.0 needs to tell Prime/Standard/Growth domestic common stock apart
from ETFs, REITs, TOKYO PRO Market and foreign shares.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime

import openpyxl

from surge.http_fetch import fetch
from surge.models import FetchResult, Provenance, ProviderCapabilities, RawSecurityRecord
from surge.normalize import normalize_jp_code
from surge.providers.base import LATENCY_DAILY_BATCH, ROLE_SECURITY_MASTER

SOURCE_ID = "jpx_listed_issues"
DEFAULT_URL = "https://www.jpx.co.jp/markets/statistics-equities/misc/tvdivq0000001vg2-att/data_j.xlsx"

# The market/product column says which segment a line is on, not what kind of
# instrument it is: preferred shares and bond-type class shares sit on Prime
# next to ordinary shares. universe-1.0.0 includes domestic COMMON stock only,
# so the instrument wording decides and the market column cannot.
#
# Measured against the official workbook (2026-09-16): 4,434 rows carry a four
# character code and are ordinary shares; the 7 rows with a five character code
# are exactly the special shares below.
_SPECIAL_SHARE_PATTERNS: tuple[tuple[str, str], ...] = (
    # (wording in the JPX security name, security type)
    ("優先出資証券", "INVESTMENT_CERTIFICATE"),
    ("優先株式", "PREFERRED"),
    ("種類株式", "OTHER"),          # 社債型種類株式 etc: not ordinary common stock
    ("新株予約権", "WARRANT"),
    ("優先出資", "INVESTMENT_CERTIFICATE"),
)

# A JPX securities code is four characters for an ordinary share. A fifth
# character that is not "0" marks a different instrument of the same issuer.
# The spec itself (証券コード協議会) is published as a PDF we have not parsed, so
# this is used only as a guard: it never promotes anything to COMMON_STOCK, it
# stops us calling something common stock when the name did not say so.
_ORDINARY_CODE_LENGTH = 4

# JPX "市場・商品区分" -> (segment code, security type)
SEGMENT_MAP: dict[str, tuple[str, str]] = {
    "プライム（内国株式）": ("PRIME_DOMESTIC", "COMMON_STOCK"),
    "スタンダード（内国株式）": ("STANDARD_DOMESTIC", "COMMON_STOCK"),
    "グロース（内国株式）": ("GROWTH_DOMESTIC", "COMMON_STOCK"),
    "プライム（外国株式）": ("PRIME_FOREIGN", "FOREIGN_COMMON_STOCK"),
    "スタンダード（外国株式）": ("STANDARD_FOREIGN", "FOREIGN_COMMON_STOCK"),
    "グロース（外国株式）": ("GROWTH_FOREIGN", "FOREIGN_COMMON_STOCK"),
    "PRO Market": ("PRO_MARKET", "COMMON_STOCK"),
    "ETF・ETN": ("ETF_ETN", "ETF"),
    "REIT・ベンチャーファンド・カントリーファンド・インフラファンド": ("REIT_FUND", "REIT_OR_FUND"),
    "出資証券": ("INVESTMENT_CERTIFICATE", "INVESTMENT_CERTIFICATE"),
}


class JpxListedIssuesProvider:
    """Reads the official JPX workbook. No credentials required."""

    provider_id = "jpx_listed_issues"

    def __init__(self, url: str = DEFAULT_URL, *, user_agent: str | None = None) -> None:
        self._url = url
        self._user_agent = user_agent

    def capabilities(self) -> ProviderCapabilities:
        return ProviderCapabilities(
            provider_id=self.provider_id,
            markets=("JP",),
            roles=(ROLE_SECURITY_MASTER,),
            update_schedule="monthly workbook published by JPX",
            latency_class=LATENCY_DAILY_BATCH,
            license_scope="public exchange publication",
            provides_delisted=False,
        )

    def fetch_security_master(self) -> FetchResult:
        kwargs = {"user_agent": self._user_agent} if self._user_agent else {}
        response = fetch(self._url, **kwargs)
        records, errors = parse_workbook(response.body, observed_at=response.received_at)

        provenance = Provenance(
            source_id=SOURCE_ID,
            endpoint=self._url,
            requested_at=response.requested_at,
            received_at=response.received_at,
            http_status=response.status,
            bytes=response.bytes,
            content_sha256=response.sha256,
            item_count=len(records),
            observed_at=response.received_at,
            available_at=response.received_at,
        )
        return FetchResult(records=tuple(records), provenance=provenance, errors=tuple(errors))


def classify_special_share(code: str, name: str) -> tuple[str, str] | None:
    """Detect an instrument that is not ordinary common stock.

    Returns (security_type, evidence) or None when the row really is an
    ordinary share. Nothing here ever returns COMMON_STOCK: the only outcomes
    are a specific non-common type or UNKNOWN, which the universe step leaves
    UNRESOLVED instead of silently including.
    """

    for wording, security_type in _SPECIAL_SHARE_PATTERNS:
        if wording in name:
            return security_type, f"name contains {wording}"

    # The name did not say, but the code shape says it is not the issuer's
    # ordinary share. Do not guess what it is.
    if len(code) > _ORDINARY_CODE_LENGTH and not code[_ORDINARY_CODE_LENGTH:].strip("0"):
        return None  # e.g. "13010": the ordinary share written in five characters
    if len(code) > _ORDINARY_CODE_LENGTH:
        return "UNKNOWN", f"code {code} is not an ordinary share code"

    return None


def parse_workbook(payload: bytes, *, observed_at: datetime | None = None) -> tuple[list[RawSecurityRecord], list[str]]:
    """Parse data_j.xlsx into raw records.

    Unknown market categories are reported as UNKNOWN rather than guessed, so the
    universe step can leave them unresolved instead of silently dropping them.
    """

    observed_at = observed_at or datetime.now(UTC)
    workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    if header is None or len(header) < 4:
        raise ValueError("unexpected JPX workbook layout")

    records: list[RawSecurityRecord] = []
    errors: list[str] = []

    for raw_row in rows:
        if raw_row is None or raw_row[1] in (None, ""):
            continue
        code = normalize_jp_code(raw_row[1])
        name = str(raw_row[2] or "").strip()
        category = str(raw_row[3] or "").strip()

        mapped = SEGMENT_MAP.get(category)
        if mapped is None:
            segment_code, security_type = "UNKNOWN", "UNKNOWN"
            errors.append(f"unmapped JPX category for {code}: {category!r}")
        else:
            segment_code, security_type = mapped

        evidence: dict[str, object] = {"jpx_category": category}
        if security_type in ("COMMON_STOCK", "FOREIGN_COMMON_STOCK"):
            special = classify_special_share(code, name)
            if special is not None:
                security_type, reason = special
                evidence["special_share"] = reason
                evidence["segment_says"] = mapped[1]

        records.append(
            RawSecurityRecord(
                source_id=SOURCE_ID,
                source_record_id=code,
                market_code="JP",
                exchange_id="XTKS",
                local_code=code,
                symbol=code,
                name=name,
                security_type=security_type,
                market_segment_code=segment_code,
                market_segment_name=category or None,
                currency="JPY",
                country="JP",
                listing_status="LISTED",
                type_evidence=evidence,
            )
        )

    return records, errors
