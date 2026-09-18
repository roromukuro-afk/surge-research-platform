"""The evaluation's JP universe: JPX's listed-issues workbook, domestic common stock.

The production adapter (``providers.jpx_listed``) reads code, name and market
category. The evaluation also needs the 33-industry sector and the size
category, which sit in the same workbook, so it reads them here rather than
changing the production parser.

These are the workbook's **current** values: a security delisted since is
absent (survivorship), and a sector or segment that changed is shown as it is
now. Phase A records both as stated limitations (jev-evaluation-design.md 9-4).
"""

from __future__ import annotations

import io
from dataclasses import dataclass

from surge.providers.jpx_listed import DEFAULT_URL, classify_special_share

#: The three domestic-stock segments. ETFs, REITs, foreign stocks and the like
#: are not in the prediction universe.
DOMESTIC_SEGMENTS = {
    "プライム（内国株式）": "PRIME",
    "スタンダード（内国株式）": "STANDARD",
    "グロース（内国株式）": "GROWTH",
}


@dataclass(frozen=True)
class ListedIssue:
    code: str
    name: str
    segment: str
    sector33: str
    size_category: str


def parse_listed_issues(payload: bytes) -> list[ListedIssue]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    rows = workbook[workbook.sheetnames[0]].iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows)]

    def column(name: str) -> int:
        for index, value in enumerate(header):
            if value == name:
                return index
        raise ValueError(f"JPX workbook has no {name!r} column (header: {header})")

    code_i, name_i, segment_i = column("コード"), column("銘柄名"), column("市場・商品区分")
    sector_i, size_i = column("33業種区分"), column("規模区分")
    issues = []
    for row in rows:
        if not row or row[code_i] in (None, ""):
            continue
        code = str(row[code_i]).strip()
        if code.endswith(".0"):
            code = code[:-2]
        name = str(row[name_i] or "").strip()
        segment = DOMESTIC_SEGMENTS.get(str(row[segment_i] or "").strip())
        if segment is None or len(code) != 4 or classify_special_share(code, name) is not None:
            continue
        issues.append(ListedIssue(
            code=code,
            name=name,
            segment=segment,
            sector33=str(row[sector_i] or "").strip(),
            size_category=str(row[size_i] or "").strip(),
        ))
    return issues


def fetch_listed_issues(url: str = DEFAULT_URL) -> tuple[list[ListedIssue], str]:
    """The workbook as JPX publishes it now, and its SHA-256."""

    from surge.http_fetch import fetch

    response = fetch(url)
    if response.status != 200:
        raise RuntimeError(f"JPX returned {response.status} for the listed-issues workbook")
    return parse_listed_issues(response.body), response.sha256


__all__ = ["DOMESTIC_SEGMENTS", "ListedIssue", "fetch_listed_issues", "parse_listed_issues"]
